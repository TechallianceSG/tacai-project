"""Document loading utilities for common recruitment input files."""

from __future__ import annotations

from pathlib import Path
import subprocess


class DocumentLoadError(RuntimeError):
    """Raised when a document cannot be loaded as text."""


class DocumentLoader:
    """Loads plain text from TXT, Markdown, PDF, and Word-style documents."""

    TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
    WORD_SUFFIXES = {".docx", ".doc", ".rtf"}

    def load(self, path: Path) -> str:
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {path}")
        if not path.is_file():
            raise DocumentLoadError(f"Document path is not a file: {path}")

        suffix = path.suffix.lower()
        if suffix in self.TEXT_SUFFIXES:
            return self._load_text(path)
        if suffix == ".pdf":
            return self._load_pdf(path)
        if suffix == ".docx":
            return self._load_docx(path)
        if suffix in {".doc", ".rtf"}:
            return self._load_with_textutil(path)

        raise DocumentLoadError(
            f"Unsupported file type '{suffix}'. Supported: .txt, .md, .pdf, .docx, .doc, .rtf"
        )

    def _load_text(self, path: Path) -> str:
        for encoding in ("utf-8", "utf-8-sig", "cp932", "shift_jis", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise DocumentLoadError(f"Unable to decode text file: {path}")

    def _load_pdf(self, path: Path) -> str:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except ImportError as exc:
            raise DocumentLoadError("PDF support requires dependency 'pypdf'. Install with: python3 -m pip install pypdf") from exc

        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(page for page in pages if page)
        if not text.strip():
            raise DocumentLoadError(f"No extractable text found in PDF: {path}")
        return text

    def _load_docx(self, path: Path) -> str:
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError as exc:
            raise DocumentLoadError("Word .docx support requires dependency 'python-docx'. Install with: python3 -m pip install python-docx") from exc

        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_cells: list[str] = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        table_cells.append(value)
        text = "\n".join(paragraphs + table_cells)
        if not text.strip():
            raise DocumentLoadError(f"No extractable text found in Word document: {path}")
        return text

    def _load_with_textutil(self, path: Path) -> str:
        try:
            result = subprocess.run(
                ["textutil", "-convert", "txt", "-stdout", str(path)],
                check=True,
                capture_output=True,
                text=True,
            )
        except (FileNotFoundError, subprocess.CalledProcessError) as exc:
            raise DocumentLoadError(
                f"Word/RTF legacy file support requires macOS textutil or conversion to .docx first: {path}"
            ) from exc
        text = result.stdout.strip()
        if not text:
            raise DocumentLoadError(f"No extractable text found in document: {path}")
        return text
