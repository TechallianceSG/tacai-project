# -*- coding: utf-8 -*-
"""Generate Bayer Shanghai AI talent market PPTX and speaker script DOCX."""
from __future__ import annotations

import html
import os
import zipfile
from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

OUT_DIR = Path('/Users/terencewang/Documents/claude-project/Bayer_AI_Talent_Sharing')
PPTX_PATH = OUT_DIR / 'Bayer上海AI人才市场专家分享.pptx'
DOCX_PATH = OUT_DIR / 'Bayer上海AI人才市场逐页话术.docx'

SLIDE_W = 12192000
SLIDE_H = 6858000
EMU_PER_INCH = 914400

COLORS = {
    'navy': '102A43',
    'blue': '1F4E79',
    'light_blue': 'D9EAF7',
    'green': '78BE20',
    'dark_green': '4E7D14',
    'gray': '52606D',
    'light_gray': 'F4F7FA',
    'white': 'FFFFFF',
    'black': '1F2933',
    'orange': 'F59E0B',
    'red': 'C2410C',
    'purple': '6D28D9',
}

FONT = 'Microsoft YaHei'


def emu(inches: float) -> int:
    return int(inches * EMU_PER_INCH)


def esc(text: str) -> str:
    return html.escape(str(text), quote=False)


def run(text: str, size: int = 1800, color: str = COLORS['black'], bold: bool = False):
    b = ' b="1"' if bold else ''
    return (
        f'<a:r><a:rPr lang="zh-CN" sz="{size}"{b}>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface="{FONT}"/>'
        f'</a:rPr><a:t>{esc(text)}</a:t></a:r>'
    )


def para(text: str, size: int = 1800, color: str = COLORS['black'], bold: bool = False,
         align: str = 'l', bullet: bool = False, level: int = 0, line_spacing: int | None = None):
    mar = 0
    indent = 0
    bullet_xml = '<a:buNone/>'
    if bullet:
        mar = 285750 + level * 285750
        indent = -171450
        bullet_xml = '<a:buChar char="•"/>'
    ln = f'<a:lnSpc><a:spcPct val="{line_spacing}"/></a:lnSpc>' if line_spacing else ''
    return (
        f'<a:p><a:pPr algn="{align}" marL="{mar}" indent="{indent}">{bullet_xml}{ln}</a:pPr>'
        f'{run(text, size=size, color=color, bold=bold)}</a:p>'
    )


def shape_xml(shape_id: int, name: str, x: int, y: int, w: int, h: int, paragraphs: list[str],
              fill: str | None = None, line: str | None = None, radius: bool = False,
              valign: str = 'top', margin: int = 110000):
    fill_xml = '<a:noFill/>' if fill is None else f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
    line_xml = '<a:ln><a:noFill/></a:ln>' if line is None else f'<a:ln w="12700"><a:solidFill><a:srgbClr val="{line}"/></a:solidFill></a:ln>'
    geom = 'roundRect' if radius else 'rect'
    anchor = {'top': 't', 'middle': 'ctr', 'bottom': 'b'}.get(valign, 't')
    return f'''
<p:sp>
  <p:nvSpPr><p:cNvPr id="{shape_id}" name="{esc(name)}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>{fill_xml}{line_xml}</p:spPr>
  <p:txBody><a:bodyPr wrap="square" anchor="{anchor}" lIns="{margin}" tIns="{margin}" rIns="{margin}" bIns="{margin}"/><a:lstStyle/>{''.join(paragraphs)}</p:txBody>
</p:sp>'''


def line_xml(shape_id: int, x1: int, y1: int, x2: int, y2: int, color: str = COLORS['green'], width: int = 25400):
    # Represent a straight connector via xfrm ext and prstGeom line.
    x = min(x1, x2); y = min(y1, y2); w = abs(x2 - x1) or 1; h = abs(y2 - y1) or 1
    flip_h = ' flipH="1"' if x2 < x1 else ''
    flip_v = ' flipV="1"' if y2 < y1 else ''
    return f'''
<p:cxnSp>
  <p:nvCxnSpPr><p:cNvPr id="{shape_id}" name="Connector {shape_id}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>
  <p:spPr><a:xfrm{flip_h}{flip_v}><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom><a:ln w="{width}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln></p:spPr>
</p:cxnSp>'''


def slide_xml(shapes: list[str], bg: str = COLORS['white']) -> str:
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:bg><p:bgPr><a:solidFill><a:srgbClr val="{bg}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>
      {''.join(shapes)}
    </p:spTree>
  </p:cSld>
  <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>'''


def title_box(sid: int, title: str, subtitle: str | None = None, dark: bool = False):
    color = COLORS['white'] if dark else COLORS['navy']
    paras = [para(title, size=3000, color=color, bold=True)]
    if subtitle:
        paras.append(para(subtitle, size=1450, color=COLORS['light_blue'] if dark else COLORS['gray']))
    return shape_xml(sid, 'Title', emu(0.45), emu(0.25), emu(12.3), emu(0.85 if not subtitle else 1.1), paras, fill=None, margin=0)


def footer(shape_id: int, text='TAC | Confidential discussion draft'):
    return shape_xml(shape_id, 'Footer', emu(0.55), emu(7.1), emu(12.25), emu(0.25),
                     [para(text, size=800, color='8A94A6'), para('2026', size=800, color='8A94A6', align='r')], fill=None, margin=0)


def make_slides():
    slides = []

    # Slide 1
    shapes = [
        shape_xml(2, 'Hero', 0, 0, SLIDE_W, SLIDE_H, [], fill=COLORS['navy']),
        shape_xml(3, 'Accent', 0, 0, emu(0.25), SLIDE_H, [], fill=COLORS['green']),
        shape_xml(4, 'Title', emu(0.75), emu(1.35), emu(11.6), emu(1.6), [
            para('上海AI人才市场与Bayer人才吸引策略', size=3400, color=COLORS['white'], bold=True),
            para('从“抢算法人才”到“赢得 AI × Life Science 复合型人才”', size=1800, color=COLORS['light_blue'])
        ], fill=None, margin=0),
        shape_xml(5, 'Meta', emu(0.78), emu(5.2), emu(8.5), emu(0.75), [
            para('面向Bayer中国HR、业务、销售、财务团队的30分钟专家分享', size=1500, color=COLORS['white']),
            para('Prepared by TAC | Recruitment & Talent Advisory', size=1200, color='B8C7D9')
        ], fill=None, margin=0),
    ]
    slides.append(slide_xml(shapes, bg=COLORS['navy']))

    # Slide 2
    shapes = [title_box(2, '今天的核心判断：AI人才不是一个市场'),
              shape_xml(3, 'Key', emu(0.75), emu(1.45), emu(11.9), emu(1.0), [
                  para('Bayer不应该把自己定位为“抢AI人才”，而应定位为“赢得AI × Life Science人才”。', size=2300, color=COLORS['white'], bold=True)
              ], fill=COLORS['blue'], radius=True, valign='middle')]
    x_positions = [0.85, 3.25, 5.65, 8.05, 10.45]
    labels = [('基础模型', '论文/模型/算力'), ('AI工程平台', 'MLOps/部署/数据'), ('业务应用', '产品化/Agent/ROI'), ('AI×行业科学', '药物/医学/农业'), ('AI治理', '合规/伦理/风险')]
    for i, (t, s) in enumerate(labels):
        shapes.append(shape_xml(10+i, f'Card{i}', emu(x_positions[i]), emu(3.0), emu(1.95), emu(1.45), [
            para(t, size=1500, color=COLORS['navy'], bold=True, align='c'),
            para(s, size=1000, color=COLORS['gray'], align='c')
        ], fill=COLORS['light_gray'], line='D5DEE8', radius=True, valign='middle'))
    shapes.append(shape_xml(20, 'Bottom', emu(0.85), emu(5.15), emu(11.45), emu(0.95), [
        para('关键问题不是“市场上有没有AI人才”，而是：Bayer到底要赢哪一类AI人才？', size=1850, color=COLORS['dark_green'], bold=True, align='c')
    ], fill='EFF8E8', line='CDEBBA', radius=True, valign='middle'))
    shapes.append(footer(30))
    slides.append(slide_xml(shapes, bg=COLORS['white']))

    # Slide 3
    shapes = [title_box(2, '上海AI人才市场的位置：产业场景密度是独特优势')]
    cities = [
        ('北京', '基础研究 / 大模型 / 高校院所', 'AI Scientist、AI for Science'),
        ('上海', '生命科学 / 金融 / 汽车 / 外企 / 企业AI', 'Bayer主战场：产业AI落地'),
        ('杭州', '云计算 / 电商 / Agent / 互联网产品', 'AI产品、平台、应用人才'),
        ('深圳', '硬件 / 机器人 / 工业AI / 供应链', '工程化、自动化、设备AI'),
        ('苏州/南京', '生物医药 / 制造 / 研发外溢', '生命科学人才补充池'),
    ]
    shapes.append(shape_xml(3, 'Header', emu(0.75), emu(1.35), emu(11.8), emu(0.55), [
        para('城市', size=1150, color=COLORS['white'], bold=True), para('AI人才特点', size=1150, color=COLORS['white'], bold=True), para('对Bayer的意义', size=1150, color=COLORS['white'], bold=True)
    ], fill=COLORS['blue']))
    y = 2.0
    for i, (c, f, m) in enumerate(cities):
        fill = 'FFFFFF' if i % 2 == 0 else 'F6F8FB'
        shapes.append(shape_xml(10+i*3, 'city', emu(0.75), emu(y), emu(1.8), emu(0.72), [para(c, size=1300, color=COLORS['navy'], bold=True, align='c')], fill=fill, line='E3E8EF', valign='middle'))
        shapes.append(shape_xml(11+i*3, 'feature', emu(2.55), emu(y), emu(5.1), emu(0.72), [para(f, size=1200, color=COLORS['black'])], fill=fill, line='E3E8EF', valign='middle'))
        shapes.append(shape_xml(12+i*3, 'meaning', emu(7.65), emu(y), emu(4.9), emu(0.72), [para(m, size=1200, color=COLORS['dark_green'], bold=(c=='上海'))], fill=fill, line='E3E8EF', valign='middle'))
        y += 0.78
    shapes.append(shape_xml(40, 'Insight', emu(0.9), emu(6.1), emu(11.2), emu(0.55), [
        para('观点：上海不是只靠“AI热度”，而是靠生命科学、金融、汽车、制造等真实产业问题吸引AI人才。', size=1300, color=COLORS['navy'], bold=True, align='c')
    ], fill='EAF5FF', line='BFDDF4', radius=True, valign='middle'))
    shapes.append(footer(41))
    slides.append(slide_xml(shapes, bg=COLORS['white']))

    # Slide 4
    shapes = [title_box(2, '上海AI人才五类分布：Bayer不要一把抓')]
    segs = [
        ('基础模型/算法研究', '大模型公司、高校、AI实验室', '少量高端引入：AI for Science'),
        ('AI工程化/MLOps', '云厂商、互联网平台、SaaS', '企业级AI平台与部署能力'),
        ('AI产品/业务应用', 'B端产品、咨询、数字化团队', '把AI做成业务工具'),
        ('AI×生命科学复合型', '生信、计算化学、医学数据、药物发现', '最值得重点争夺'),
        ('AI治理/合规/风险', '隐私、法务、质量、模型风险', '跨国药企差异化优势'),
    ]
    colors = [COLORS['blue'], COLORS['purple'], COLORS['green'], COLORS['orange'], COLORS['gray']]
    y = 1.35
    for i, (t, src, val) in enumerate(segs):
        shapes.append(shape_xml(10+i, 'segment', emu(0.75), emu(y), emu(11.8), emu(0.85), [
            para(t, size=1350, color=COLORS['white'], bold=True),
            para(f'来源：{src}  |  Bayer意义：{val}', size=1050, color=COLORS['white'])
        ], fill=colors[i], radius=True, valign='middle'))
        y += 1.0
    shapes.append(shape_xml(30, 'Note', emu(0.9), emu(6.25), emu(11.2), emu(0.45), [para('建议聚焦：AI×生命科学、AI工程平台、AI业务转译三类人才。', size=1250, color=COLORS['navy'], bold=True, align='c')], fill='F6FAFF', line='CFE2F3', radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 5
    shapes = [title_box(2, '人才流动：上海AI人才为什么跳槽？')]
    trends = [
        ('从互联网/大厂 → 产业AI', '不只优化点击率，希望解决更真实、更长期的问题'),
        ('从纯算法 → 产品化/Agent化/业务落地', '关注能否上线、能否产生ROI，而非只做Demo'),
        ('从外企 → 本土科技/创业公司', '担心决策慢、技术栈旧、数据拿不到、岗位离业务远'),
        ('从高薪机会 → 高确定性机会', '经历市场波动后，重新重视稳定平台、长期问题、品牌信任'),
    ]
    for i, (t, d) in enumerate(trends):
        x = 0.8 + (i % 2) * 6.0
        y = 1.55 + (i // 2) * 2.05
        shapes.append(shape_xml(10+i, 'trend', emu(x), emu(y), emu(5.35), emu(1.45), [
            para(t, size=1500, color=COLORS['navy'], bold=True),
            para(d, size=1120, color=COLORS['gray'])
        ], fill=COLORS['light_gray'], line='DDE5ED', radius=True))
    shapes.append(shape_xml(20, 'Implication', emu(0.85), emu(5.75), emu(11.55), emu(0.65), [para('对Bayer的机会：把“生命科学真实问题 + 全球平台 + 中国落地”讲成具体职业机会。', size=1400, color=COLORS['white'], bold=True, align='c')], fill=COLORS['blue'], radius=True, valign='middle'))
    shapes.append(footer(21))
    slides.append(slide_xml(shapes))

    # Slide 6
    shapes = [title_box(2, '候选人真正关心什么：不是只问薪资')]
    questions = [
        ('数据在哪里？', '能否接触真实、合规、可用的数据'),
        ('业务owner是谁？', '业务是否真正投入，而不是“让AI团队试试看”'),
        ('能不能上线？', '项目是否进入流程、触达用户、产生业务结果'),
        ('技术路线是否现代？', '工具链、架构、MLOps、GenAI能力是否跟得上'),
        ('职业路径是什么？', '是否有专家路线，而不只是转管理'),
    ]
    y = 1.35
    for i, (q, a) in enumerate(questions):
        shapes.append(shape_xml(10+i, 'q', emu(0.9), emu(y), emu(3.0), emu(0.72), [para(q, size=1300, color=COLORS['white'], bold=True, align='c')], fill=COLORS['green'], radius=True, valign='middle'))
        shapes.append(shape_xml(20+i, 'a', emu(4.05), emu(y), emu(8.0), emu(0.72), [para(a, size=1250, color=COLORS['black'])], fill='FFFFFF', line='DDE5ED', radius=True, valign='middle'))
        y += 0.92
    shapes.append(shape_xml(40, 'Insight', emu(0.95), emu(6.05), emu(11.1), emu(0.55), [para('金句：AI人才不怕复杂问题，怕的是复杂组织。', size=1500, color=COLORS['navy'], bold=True, align='c')], fill='EFF8E8', line='CDEBBA', radius=True, valign='middle'))
    shapes.append(footer(41))
    slides.append(slide_xml(shapes))

    # Slide 7
    shapes = [title_box(2, 'Bayer的Selling Points：不只是外企品牌')]
    points = [
        ('生命科学真实场景', 'AI创业公司很难拥有的真实问题、专业数据和长期场景'),
        ('全球平台 + 中国场景', '中国市场复杂度与全球生命科学网络结合'),
        ('合规与信任', '医疗/生命科学AI落地需要被信任、被合规使用'),
        ('多业务线应用广度', '医药、消费者健康、农业、供应链、财务、HR等多场景'),
    ]
    for i, (t, d) in enumerate(points):
        x = 0.8 + (i % 2) * 6.0
        y = 1.55 + (i // 2) * 2.0
        shapes.append(shape_xml(10+i, 'point', emu(x), emu(y), emu(5.4), emu(1.45), [
            para(t, size=1500, color=COLORS['navy'], bold=True),
            para(d, size=1120, color=COLORS['gray'])
        ], fill='F7FBFF', line='CFE2F3', radius=True))
    shapes.append(shape_xml(20, 'Quote', emu(0.85), emu(5.75), emu(11.5), emu(0.7), [para('关键：把“百年品牌”翻译成“AI人才加入后能解决什么具体问题”。', size=1450, color=COLORS['white'], bold=True, align='c')], fill=COLORS['navy'], radius=True, valign='middle'))
    shapes.append(footer(21))
    slides.append(slide_xml(shapes))

    # Slide 8
    shapes = [title_box(2, '把Bayer优势翻译成AI候选人语言')]
    rows = [
        ('传统表达', '候选人更想听到的表达'),
        ('我们是全球领先生命科学公司', '你会参与真实生命科学问题，而不是只做流量优化'),
        ('我们有全球平台', '你能和中国业务、全球专家、跨区域团队一起把AI落地'),
        ('我们重视合规', '你的AI方案有机会进入被信任、可审计、可规模化的业务流程'),
        ('我们有多元业务', '你可以在研发、医学、商业、供应链、财务等多场景放大能力'),
    ]
    y = 1.25
    for i, (l, r) in enumerate(rows):
        is_head = i == 0
        fill = COLORS['blue'] if is_head else ('FFFFFF' if i % 2 else 'F6F8FB')
        color = COLORS['white'] if is_head else COLORS['black']
        shapes.append(shape_xml(10+i*2, 'left', emu(0.85), emu(y), emu(4.0), emu(0.82), [para(l, size=1200 if not is_head else 1250, color=color, bold=is_head)], fill=fill, line='DDE5ED', valign='middle'))
        shapes.append(shape_xml(11+i*2, 'right', emu(4.85), emu(y), emu(7.2), emu(0.82), [para(r, size=1200 if not is_head else 1250, color=color, bold=is_head)], fill=fill, line='DDE5ED', valign='middle'))
        y += 0.87
    shapes.append(shape_xml(30, 'Insight', emu(0.95), emu(6.05), emu(11.1), emu(0.55), [para('EVP不是口号，而是候选人对“我加入后第一年能做成什么”的判断。', size=1350, color=COLORS['dark_green'], bold=True, align='c')], fill='EFF8E8', line='CDEBBA', radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 9
    shapes = [title_box(2, 'Bayer吸引AI人才的潜在短板：不是品牌弱，而是机会叙事不够具体')]
    gaps = [
        ('AI雇主品牌不够具体', '候选人不清楚Bayer中国AI项目、话语权、专家路线'),
        ('决策速度弱于AI-native公司', '面试、offer、数据审批、项目启动速度影响判断'),
        ('薪酬benchmark容易看错', 'AI人才对标大模型、云、金融科技、机器人，而非传统药企'),
        ('岗位设计语言过于传统', 'Data Analyst/IT Manager式JD会被理解为BI或传统IT岗位'),
        ('AI团队与业务团队脱节', '没有业务owner、数据、上线路径，候选人会担心只做PoC'),
    ]
    y = 1.25
    for i, (g, d) in enumerate(gaps):
        shapes.append(shape_xml(10+i, 'gap', emu(0.75), emu(y), emu(11.8), emu(0.78), [
            para(f'{i+1}. {g}', size=1300, color=COLORS['navy'], bold=True),
            para(d, size=1050, color=COLORS['gray'])
        ], fill='FFF8ED' if i % 2 == 0 else 'FFFFFF', line='F4D6A0', radius=True, valign='middle'))
        y += 0.9
    shapes.append(shape_xml(30, 'Quote', emu(0.95), emu(6.05), emu(11.1), emu(0.55), [para('建议：AI关键岗位设置fast-track hiring，业务高层参与sell，面试周期控制在2-3周。', size=1250, color=COLORS['white'], bold=True, align='c')], fill=COLORS['orange'], radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 10
    shapes = [title_box(2, 'Bayer最该争夺的三类AI人才')]
    archetypes = [
        ('1', 'AI × Life Science复合型人才', 'AI药物发现、计算生物、生物信息、医学数据、农业科技AI', '用科学问题、真实数据、全球专家网络吸引'),
        ('2', '企业级AI产品与平台人才', '云厂商、SaaS、大厂AI平台、企业数字化团队', '用多场景、从0到1、平台影响力吸引'),
        ('3', 'AI Business Translator', '懂业务流程、数据逻辑、AI边界、ROI、stakeholder管理', '把业务语言翻译成AI问题，把AI能力翻译成业务结果'),
    ]
    for i, (num, t, bg, hook) in enumerate(archetypes):
        x = 0.75 + i * 4.05
        shapes.append(shape_xml(10+i, 'num', emu(x), emu(1.45), emu(0.55), emu(0.55), [para(num, size=1500, color=COLORS['white'], bold=True, align='c')], fill=COLORS['green'], radius=True, valign='middle'))
        shapes.append(shape_xml(20+i, 'arch', emu(x), emu(2.05), emu(3.65), emu(3.2), [
            para(t, size=1400, color=COLORS['navy'], bold=True),
            para('目标背景：', size=950, color=COLORS['gray'], bold=True),
            para(bg, size=980, color=COLORS['gray']),
            para('吸引方式：', size=950, color=COLORS['gray'], bold=True),
            para(hook, size=980, color=COLORS['dark_green'])
        ], fill='F7FBFF', line='CFE2F3', radius=True))
    shapes.append(shape_xml(40, 'Insight', emu(0.9), emu(5.85), emu(11.25), emu(0.6), [para('观点：Bayer最稀缺的可能不是第100个算法工程师，而是能把AI带进业务流程的人。', size=1300, color=COLORS['navy'], bold=True, align='c')], fill='EFF8E8', line='CDEBBA', radius=True, valign='middle'))
    shapes.append(footer(41))
    slides.append(slide_xml(shapes))

    # Slide 11
    shapes = [title_box(2, 'AI人才吸引力公式：HR、业务、财务可以共同使用')]
    shapes.append(shape_xml(3, 'Formula', emu(0.9), emu(1.45), emu(11.2), emu(1.15), [
        para('AI人才吸引力 = 问题价值 × 数据/场景质量 × 决策速度 × 职业成长 ÷ 组织摩擦', size=2100, color=COLORS['white'], bold=True, align='c')
    ], fill=COLORS['navy'], radius=True, valign='middle'))
    factors = [
        ('问题价值', '我做的事情是否有真实意义？'),
        ('数据/场景质量', '有没有真实数据和真实用户？'),
        ('决策速度', '我进去以后推得动吗？'),
        ('职业成长', '我是管理者，还是也能成为技术专家？'),
        ('组织摩擦', '审批、合规、跨部门是否会消耗热情？'),
    ]
    for i, (f, q) in enumerate(factors):
        x = 0.75 + i * 2.35
        shapes.append(shape_xml(10+i, 'factor', emu(x), emu(3.25), emu(2.05), emu(1.6), [
            para(f, size=1250, color=COLORS['navy'], bold=True, align='c'),
            para(q, size=900, color=COLORS['gray'], align='c')
        ], fill=COLORS['light_gray'], line='DDE5ED', radius=True, valign='middle'))
    shapes.append(shape_xml(30, 'Note', emu(0.9), emu(5.85), emu(11.2), emu(0.55), [para('财务视角：关键AI岗位不是普通headcount成本，而是未来能力建设投资。', size=1350, color=COLORS['white'], bold=True, align='c')], fill=COLORS['blue'], radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 12
    shapes = [title_box(2, 'HR、业务、销售、财务分别该做什么')]
    actions = [
        ('HR', '建立AI岗位族群、人才地图、专属candidate pitch、fast-track流程'),
        ('业务部门', '定义真实场景、投入业务owner、承诺上线路径和成功指标'),
        ('销售/商业', '提出客户分层、销售赋能、市场洞察等可量化AI场景'),
        ('财务', '区分普通数据岗位与战略AI岗位，建立差异化预算逻辑'),
    ]
    y = 1.4
    for i, (role, act) in enumerate(actions):
        color = [COLORS['blue'], COLORS['green'], COLORS['orange'], COLORS['purple']][i]
        shapes.append(shape_xml(10+i*2, 'role', emu(0.85), emu(y), emu(2.0), emu(0.85), [para(role, size=1500, color=COLORS['white'], bold=True, align='c')], fill=color, radius=True, valign='middle'))
        shapes.append(shape_xml(11+i*2, 'act', emu(3.05), emu(y), emu(8.85), emu(0.85), [para(act, size=1250, color=COLORS['black'])], fill='FFFFFF', line='DDE5ED', radius=True, valign='middle'))
        y += 1.1
    shapes.append(shape_xml(30, 'Insight', emu(0.95), emu(6.05), emu(11.1), emu(0.55), [para('AI招聘不是HR单点动作，而是业务战略、组织速度和投资逻辑的共同结果。', size=1350, color=COLORS['navy'], bold=True, align='c')], fill='EFF8E8', line='CDEBBA', radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 13
    shapes = [title_box(2, '90天行动建议：从人才地图到标杆招聘案例')]
    phases = [
        ('0-30天｜明确人才战场', ['定义关键AI岗位族群', '区分普通数据岗与战略AI岗', '完成上海/北京/杭州/苏州人才地图', '梳理真实AI业务场景']),
        ('31-60天｜重塑候选人叙事', ['制作关键岗位AI Candidate Pitch', '训练HR与业务leader讲岗位故事', '优化JD语言与岗位名称', '建立AI岗位快速面试流程']),
        ('61-90天｜建立标杆案例', ['选2-3个关键岗位高质量招聘', '业务高层参与候选人吸引', '完成市场薪酬校准', '形成Bayer中国AI人才吸引playbook']),
    ]
    for i, (phase, items) in enumerate(phases):
        x = 0.75 + i * 4.05
        paras = [para(phase, size=1300, color=COLORS['white'], bold=True)] + [para(item, size=950, color=COLORS['black'], bullet=True) for item in items]
        shapes.append(shape_xml(10+i, 'phase', emu(x), emu(1.45), emu(3.65), emu(4.55), paras, fill='F7FBFF', line='CFE2F3', radius=True))
        shapes.append(shape_xml(20+i, 'phasehead', emu(x), emu(1.45), emu(3.65), emu(0.65), [para(phase, size=1200, color=COLORS['white'], bold=True, align='c')], fill=[COLORS['blue'], COLORS['green'], COLORS['orange']][i], radius=True, valign='middle'))
    shapes.append(footer(31))
    slides.append(slide_xml(shapes))

    # Slide 14
    shapes = [title_box(2, '结尾三句话：Bayer在上海AI人才市场并非弱势方')]
    takeaways = [
        ('1', '上海AI人才市场已从“热闹”进入“分层”', '关键不是所有AI人才都要抢，而是要锁定Bayer最该赢的人才。'),
        ('2', 'Bayer的优势不是速度和薪资，而是场景、平台、合规和长期问题', '这些优势要被翻译成AI人才听得懂的岗位机会。'),
        ('3', '未来要赢的是“AI业务落地人才战”', '最关键的人，往往是能把AI带进医学、销售、研发和运营流程的人。'),
    ]
    y = 1.45
    for num, t, d in takeaways:
        shapes.append(shape_xml(10+int(num), 'num', emu(0.9), emu(y), emu(0.7), emu(0.7), [para(num, size=1700, color=COLORS['white'], bold=True, align='c')], fill=COLORS['green'], radius=True, valign='middle'))
        shapes.append(shape_xml(20+int(num), 'take', emu(1.85), emu(y), emu(10.0), emu(0.9), [para(t, size=1450, color=COLORS['navy'], bold=True), para(d, size=1050, color=COLORS['gray'])], fill='FFFFFF', line='DDE5ED', radius=True, valign='middle'))
        y += 1.35
    shapes.append(shape_xml(40, 'Closing', emu(0.9), emu(5.85), emu(11.2), emu(0.75), [para('如果Bayer能把自己表达为“AI × Life Science的真实应用平台”，它在上海有非常独特的位置。', size=1450, color=COLORS['white'], bold=True, align='c')], fill=COLORS['navy'], radius=True, valign='middle'))
    shapes.append(footer(41))
    slides.append(slide_xml(shapes))

    # Slide 15
    shapes = [title_box(2, 'Q&A / 参考来源')]
    sources = [
        'Shanghai Government：上海人才吸引力、AI产业与外商投资指南',
        'China Daily / Shanghai Government：上海AI产业、大模型与AI人才公开报道',
        'China.org.cn / Liepin引用：2025中国AI人才供需与招聘热度',
        'Shanghai Pudong / Zhangjiang：生物医药人才与AI交叉需求',
        'Bayer China / Bayer Global：Career、EVP、AI innovation、life sciences incubator',
    ]
    shapes.append(shape_xml(3, 'Thank', emu(0.9), emu(1.35), emu(11.2), emu(1.0), [para('谢谢，期待讨论：Bayer中国如何把AI人才战略变成业务增长能力？', size=2200, color=COLORS['white'], bold=True, align='c')], fill=COLORS['blue'], radius=True, valign='middle'))
    shapes.append(shape_xml(4, 'Sources', emu(0.95), emu(2.75), emu(11.0), emu(3.6), [para('参考来源，可用于PPT脚注或会后版本：', size=1300, color=COLORS['navy'], bold=True)] + [para(s, size=1050, color=COLORS['gray'], bullet=True) for s in sources], fill='F7FBFF', line='CFE2F3', radius=True))
    shapes.append(footer(5))
    slides.append(slide_xml(shapes))

    return slides


def content_types(n):
    overrides = ''.join(f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>' for i in range(1, n+1))
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
  {overrides}
</Types>'''


def root_rels():
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>'''


def presentation_xml(n):
    ids = ''.join(f'<p:sldId id="{255+i}" r:id="rId{i}"/>' for i in range(1, n+1))
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" saveSubsetFonts="1">
  <p:sldIdLst>{ids}</p:sldIdLst>
  <p:sldSz cx="{SLIDE_W}" cy="{SLIDE_H}" type="wide"/>
  <p:notesSz cx="6858000" cy="9144000"/>
  <p:defaultTextStyle><a:defPPr><a:defRPr lang="zh-CN"><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/></a:defRPr></a:defPPr></p:defaultTextStyle>
</p:presentation>'''


def presentation_rels(n):
    rels = ''.join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>' for i in range(1, n+1))
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>'''


def props_core():
    now = datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ')
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>Bayer上海AI人才市场专家分享</dc:title>
  <dc:creator>TAC / Claude Code</dc:creator>
  <cp:lastModifiedBy>TAC / Claude Code</cp:lastModifiedBy>
  <dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>
  <dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>
</cp:coreProperties>'''


def props_app(n):
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Application>Claude Code</Application><PresentationFormat>Widescreen</PresentationFormat><Slides>{n}</Slides></Properties>'''


def write_pptx():
    slides = make_slides()
    with zipfile.ZipFile(PPTX_PATH, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types(len(slides)))
        z.writestr('_rels/.rels', root_rels())
        z.writestr('ppt/presentation.xml', presentation_xml(len(slides)))
        z.writestr('ppt/_rels/presentation.xml.rels', presentation_rels(len(slides)))
        z.writestr('docProps/core.xml', props_core())
        z.writestr('docProps/app.xml', props_app(len(slides)))
        for i, s in enumerate(slides, 1):
            z.writestr(f'ppt/slides/slide{i}.xml', s)


scripts = [
    ('第1页｜标题页', [
        '各位Bayer的HR、业务、销售和财务同事，大家好。今天我想用猎头和人才顾问的视角，谈一谈上海AI人才市场，以及Bayer在这个市场里如何更有效地吸引真正有价值的人。',
        '我今天不会把重点放在泛泛的AI热度或宏观概念上，而是围绕一个更实用的问题：Bayer到底该争夺哪类AI人才？以及如何把Bayer自己的优势讲成候选人真正愿意听、愿意相信的机会。',
        '我的核心判断是：Bayer不应该和所有AI公司抢同一批人，而应该把战场聚焦在“AI × Life Science”复合型人才和企业级AI落地人才上。'
    ]),
    ('第2页｜核心判断：AI人才不是一个市场', [
        '过去两年，几乎所有公司都说自己需要AI人才。但从招聘市场看，AI人才并不是一个统一市场，而是至少分成五个子市场。',
        '第一类是基础模型和算法研究人才，他们关注论文、模型、算力和技术影响力。第二类是AI工程平台人才，关注MLOps、模型部署和数据平台。第三类是业务应用人才，负责把AI做成产品和流程。第四类是AI与行业科学结合的人才，比如药物发现、医学数据、农业科技。第五类是AI治理、合规和风险人才。',
        'Bayer真正要思考的是：我们是否需要每一类都去抢？我的建议是否定的。Bayer最有优势的战场，是AI与生命科学结合，以及企业级AI在真实业务里的落地。'
    ]),
    ('第3页｜上海AI人才市场的位置', [
        '中国AI人才不是平均分布的。北京更强在基础研究、大模型、高校和研究院；杭州更强在云计算、电商、互联网产品和Agent应用；深圳更强在硬件、机器人、供应链和工业AI。',
        '上海的独特性在于产业场景密度：生命科学、金融、汽车、制造、跨国公司总部和企业级数字化场景都非常集中。对Bayer来说，上海不是简单的AI人才池，而是最贴近自身业务场景的主战场。',
        '所以Bayer在上海要讲的不是“我们也做AI”，而是“我们有足够复杂、足够真实、足够长期的生命科学和商业场景，让AI人才可以产生真实影响”。'
    ]),
    ('第4页｜上海AI人才五类分布', [
        '如果把上海AI人才拆开看，大致可以分为五类。第一类是基础模型和算法研究型人才，这类人才很贵，也很挑剔，Bayer只需要少量高端引入。',
        '第二类是AI工程化和MLOps人才，这类人对Bayer非常重要，因为企业AI不是做Demo，而是要安全、稳定、合规地部署。第三类是AI产品和业务应用人才，他们能帮助销售、市场、医学、供应链、财务等部门把AI真正用起来。',
        '第四类是AI×生命科学复合型人才，这是Bayer最值得重点争夺的人群。第五类是AI治理和合规人才，这恰恰是跨国生命科学公司相比创业公司更有优势的领域。'
    ]),
    ('第5页｜人才流动：为什么跳槽', [
        '现在上海AI人才的流动有几个明显方向。第一，从互联网和大厂流向产业AI。很多人开始厌倦只做流量、推荐、广告优化，希望解决更真实、更长期的问题。',
        '第二，从纯算法流向产品化、Agent化和业务落地。大家更关心模型能不能上线、能不能创造ROI，而不是只停留在PoC。第三，有一部分人从外企流向本土科技公司或创业公司，原因往往不是不认可外企，而是担心决策慢、数据拿不到、技术栈不够新。',
        '第四，经历市场波动后，也有不少中高端人才重新重视确定性、长期平台、品牌信任和工作生活平衡。这是Bayer的机会。'
    ]),
    ('第6页｜候选人真正关心什么', [
        'AI候选人当然关心薪资，但高端候选人不会只看薪资。他们通常会问五个问题。',
        '第一，数据在哪里？没有真实、合规、可用的数据，AI项目很难有价值。第二，业务owner是谁？如果只是IT或数字化部门在推动，而业务部门没有真正投入，候选人会担心项目落不了地。第三，能不能上线？AI人才不想一直做Demo。',
        '第四，技术路线是否现代？他们会判断公司的工具链、架构、平台能力是否跟得上。第五，职业路径是什么？很多高级技术人才不一定想转管理，他们希望有专家路线。总结一句话：AI人才不怕复杂问题，怕的是复杂组织。'
    ]),
    ('第7页｜Bayer的Selling Points', [
        'Bayer对AI人才真正有吸引力的地方，不只是“外企品牌”或“百年公司”。第一是生命科学真实场景。AI创业公司可能很快，但未必拥有Bayer这样长期、复杂、真实的生命科学问题。',
        '第二是全球平台加中国场景。上海AI人才不缺本地机会，但能把中国复杂市场和全球生命科学网络结合起来的公司并不多。第三是合规和信任。AI进入医药、健康和农业后，问题不只是模型准不准，而是能否被信任、被审计、被合规使用。',
        '第四是多业务线应用广度。Bayer不是单一业务，它有医药、消费者健康、农业、供应链、商业、财务、HR等多场景，这对AI产品和业务落地人才很有吸引力。'
    ]),
    ('第8页｜把优势翻译成候选人语言', [
        '很多公司在招聘时会说“我们是全球领先企业”“我们有全球平台”“我们重视合规”。这些话没有错，但对AI候选人来说还不够具体。',
        '候选人更想听到的是：我加入后会解决什么问题？会接触什么数据？会和谁合作？我的方案能否进入真实业务流程？我第一年如何被衡量？',
        '所以Bayer的雇主品牌需要从公司介绍升级为机会叙事。比如，不只是说“我们是生命科学公司”，而是说“你会参与真实生命科学问题，而不是只做流量优化”。不只是说“我们重视合规”，而是说“你的AI方案有机会进入可审计、可信任、可规模化的业务流程”。'
    ]),
    ('第9页｜潜在短板', [
        '我也想直接讲Bayer在吸引AI人才时可能遇到的短板。第一，AI雇主品牌可能还不够具体。很多候选人知道Bayer，但不知道Bayer中国AI团队具体在做什么。',
        '第二，决策速度可能弱于AI-native公司。AI候选人会通过面试速度、offer速度和项目启动速度判断公司是否真的重视AI。第三，薪酬benchmark容易看错。AI人才的机会成本不是传统药企岗位，而是大模型公司、云厂商、金融科技、机器人和AI创业公司。',
        '第四，岗位设计语言可能过于传统。如果JD还是Data Analyst、IT Manager或Reporting Specialist的语言，高端AI候选人会认为这不是AI机会。第五，如果AI团队和业务团队脱节，候选人会担心进去后只是做PoC。'
    ]),
    ('第10页｜最该争夺的三类AI人才', [
        '我建议Bayer优先争夺三类人。第一类是AI×Life Science复合型人才，包括AI药物发现、计算生物、生物信息、医学数据和农业科技AI。这类人数量不多，但和Bayer高度匹配。',
        '第二类是企业级AI产品与平台人才，来自云厂商、SaaS、大厂AI平台和企业数字化团队。他们能帮助Bayer把AI从项目变成平台能力。',
        '第三类非常关键，叫AI Business Translator。他们未必是最强算法专家，但懂业务流程、数据逻辑、AI能力边界、ROI和stakeholder管理。对Bayer这样的大型组织来说，这类人能把业务语言翻译成AI问题，也能把AI能力翻译成业务结果。'
    ]),
    ('第11页｜AI人才吸引力公式', [
        '我建议Bayer内部可以用这个公式来判断一个岗位对AI人才的吸引力：AI人才吸引力等于问题价值，乘以数据和场景质量，乘以决策速度，乘以职业成长，再除以组织摩擦。',
        '问题价值决定候选人是否觉得这件事有意义。数据和场景质量决定AI能否真正做出结果。决策速度决定候选人是否相信自己进来后推得动。职业成长决定候选人是否愿意长期留下。组织摩擦越大，吸引力越低。',
        '这张公式特别适合HR、业务和财务一起看。因为AI招聘不是单纯薪资问题，而是岗位设计、业务投入、组织机制和投资逻辑共同决定的。'
    ]),
    ('第12页｜各部门该做什么', [
        'AI招聘不能只交给HR。HR当然很重要，需要建立AI岗位族群、人才地图、candidate pitch和快速面试流程。',
        '但业务部门同样关键。业务必须定义真实场景，明确业务owner，承诺上线路径和成功指标。销售和商业团队也要把AI场景具体化，比如客户分层、销售效率、市场洞察、预测和客户互动质量提升。',
        '财务也要参与，因为关键AI岗位不能完全按照传统岗位等级和成本逻辑处理。要区分普通数据岗位和战略AI岗位。真正关键的人才，薪酬不是单纯成本，而是能力建设投资。'
    ]),
    ('第13页｜90天行动建议', [
        '我建议Bayer可以用90天做一个小而实的动作。前30天，先明确人才战场：定义关键AI岗位族群，区分普通数据岗和战略AI岗，完成上海、北京、杭州、苏州的人才地图，并梳理内部真实AI场景。',
        '31到60天，重塑候选人叙事：为关键岗位制作AI Candidate Pitch，训练HR和业务leader讲岗位故事，优化JD语言和岗位名称，并建立AI岗位快速面试流程。',
        '61到90天，建立标杆招聘案例：选择2到3个关键岗位做高质量招聘，让业务高层参与候选人吸引，完成薪酬校准，最终形成Bayer中国AI人才吸引playbook。'
    ]),
    ('第14页｜结尾三句话', [
        '最后我想用三句话总结。第一，上海AI人才市场已经从热闹进入分层。Bayer不需要和所有AI公司抢所有人，而是要锁定最适合自己的细分人才。',
        '第二，Bayer的优势不是速度和薪资，而是生命科学场景、全球平台、合规信任和长期问题。但这些优势必须被翻译成AI人才听得懂的职业机会。',
        '第三，未来Bayer要赢的不是泛泛的AI招聘战，而是AI业务落地人才战。最关键的人，不一定是最会写模型的人，而是能把AI真正带进医学、销售、研发和运营流程的人。',
        '如果Bayer能把自己表达为“AI×Life Science的真实应用平台”，它在上海AI人才市场并不是弱势方，反而有非常独特的位置。'
    ]),
    ('第15页｜Q&A / 参考来源', [
        '以上就是今天的分享。我的建议是，会后可以围绕三个问题继续讨论：第一，Bayer中国未来12个月最关键的AI场景是什么？第二，这些场景真正需要哪类AI人才？第三，我们现在的岗位设计、薪酬逻辑和招聘流程，是否足够匹配这类人才？',
        '参考来源包括上海政府、China Daily、China.org.cn、浦东/张江生物医药人才公开信息，以及Bayer中国和Bayer全球的职业、AI创新和生命科学相关公开信息。',
        '谢谢大家，期待后续讨论。'
    ]),
]

references = [
    ('Shanghai tops China\'s talent inflow ranking', 'https://english.shanghai.gov.cn/en-Latest-TalentsinShanghai/20250513/a832a348e6144f469d3867dba179c636.html'),
    ('Artificial intelligence – Shanghai Government / 2025 Shanghai Foreign Investment Guide', 'https://english.shanghai.gov.cn/en-KeyIndustries/20250709/f4f1fb9abd8149b5b286bb87b327c1ea.html'),
    ('China’s rapid AI growth sparks hiring boom as demand outpaces supply', 'https://www.china.org.cn/business/2025-04/01/content_117799252.htm'),
    ('Pudong’s high-tech park attracts global biomedical professionals', 'https://english.shanghai.gov.cn/en-Latest-TalentsinShanghai/20251209/cac51649e1f448d095aa755cfcea4296.html'),
    ('Bayer China Career', 'https://www.bayer.com.cn/en/career'),
    ('Bayer Global: Be You. Be Bayer.', 'https://www.bayer.com/en/working-at-bayer'),
    ('Bayer Global: Unleashing the Potential of AI', 'https://www.bayer.com/en/innovation/unleashing-the-potential-of-ai'),
    ('Bayer opens life sciences incubator in Shanghai', 'https://english.shanghai.gov.cn/en-Latest-WhatsNew/20241010/962d41274cfd42629e3f40e68e5487ce.html'),
]


def set_run_font(run, font_name=FONT, size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_heading('', level=level)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=COLORS['navy'])
    return p


def write_docx():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Pt(54)
    section.right_margin = Pt(54)
    section.top_margin = Pt(54)
    section.bottom_margin = Pt(54)

    styles = doc.styles
    styles['Normal'].font.name = FONT
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    styles['Normal'].font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Bayer上海AI人才市场专家分享\n逐页演讲话术')
    set_run_font(r, size=20, bold=True, color=COLORS['navy'])
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run('30分钟版本｜面向HR、业务、销售、财务团队')
    set_run_font(r2, size=12, color=COLORS['gray'])

    add_heading(doc, '使用建议', level=1)
    for item in [
        '整体建议控制在30分钟：前14页约27分钟，最后Q&A约3分钟。',
        '现场可根据听众反应适度压缩第3-4页市场部分，把更多时间留给第7-13页Bayer策略建议。',
        '如Bayer内部业务负责人较多，重点强调“业务owner、上线路径、ROI和AI Business Translator”。',
    ]:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles['Normal']
        p.paragraph_format.left_indent = Pt(12)
        r = p.add_run('• ' + item)
        set_run_font(r)

    for title, paras in scripts:
        add_heading(doc, title, level=1)
        for text in paras:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, size=10.5, color=COLORS['black'])

    add_heading(doc, '可穿插使用的金句', level=1)
    quotes = [
        '上海AI人才不是缺供给，而是缺真正懂行业落地的人。',
        'Bayer不应该和大模型公司抢同一批人，而应该抢AI × Life Science的人。',
        'AI人才不怕复杂问题，怕的是复杂组织。',
        '外企吸引AI人才的核心，不是稳定，而是高质量问题和长期投入。',
        'Bayer最强的selling point不是“百年品牌”，而是别人没有的生命科学真实场景。',
        '未来最稀缺的不是纯算法人才，而是AI Business Translator。',
        '候选人判断一个AI岗位，会看三个问题：数据在哪里，业务owner是谁，项目能不能上线。',
        '如果AI岗位挂在传统IT语言下面，高端候选人会自动降级理解这个机会。',
        '对AI人才来说，薪资是门票，场景和速度才是决定因素。',
        'Bayer要把“合规”从限制条件，转化为AI时代的信任优势。',
    ]
    for q in quotes:
        p = doc.add_paragraph()
        r = p.add_run('• ' + q)
        set_run_font(r, size=10.5, color=COLORS['dark_green'])

    add_heading(doc, '参考来源', level=1)
    for name, url in references:
        p = doc.add_paragraph()
        r = p.add_run(f'{name}: {url}')
        set_run_font(r, size=9.5, color=COLORS['gray'])

    doc.save(DOCX_PATH)


if __name__ == '__main__':
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_pptx()
    write_docx()
    print(PPTX_PATH)
    print(DOCX_PATH)
